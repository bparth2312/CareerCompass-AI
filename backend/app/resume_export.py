from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_margins(
    cell,
    top: int = 80,
    start: int = 80,
    bottom: int = 80,
    end: int = 80
) -> None:
    """
    Set table-cell margins using Word XML values.
    """

    cell_properties = (
        cell._tc.get_or_add_tcPr()
    )

    cell_margins = (
        cell_properties.first_child_found_in(
            "w:tcMar"
        )
    )

    if cell_margins is None:
        cell_margins = OxmlElement(
            "w:tcMar"
        )

        cell_properties.append(
            cell_margins
        )

    margin_values = {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end
    }

    for margin_name, margin_value in (
        margin_values.items()
    ):
        margin_element = (
            cell_margins.find(
                qn(
                    f"w:{margin_name}"
                )
            )
        )

        if margin_element is None:
            margin_element = OxmlElement(
                f"w:{margin_name}"
            )

            cell_margins.append(
                margin_element
            )

        margin_element.set(
            qn("w:w"),
            str(margin_value)
        )

        margin_element.set(
            qn("w:type"),
            "dxa"
        )


def add_bottom_border(
    paragraph,
    color: str = "0F766E",
    size: str = "8",
    space: str = "3"
) -> None:
    """
    Add a bottom border below a paragraph.
    """

    paragraph_properties = (
        paragraph._p.get_or_add_pPr()
    )

    borders = (
        paragraph_properties.find(
            qn("w:pBdr")
        )
    )

    if borders is None:
        borders = OxmlElement(
            "w:pBdr"
        )

        paragraph_properties.append(
            borders
        )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single"
    )

    bottom.set(
        qn("w:sz"),
        size
    )

    bottom.set(
        qn("w:space"),
        space
    )

    bottom.set(
        qn("w:color"),
        color
    )

    borders.append(
        bottom
    )


def apply_document_margins(
    document: Document
) -> None:
    for section in document.sections:
        section.top_margin = Inches(
            0.55
        )

        section.bottom_margin = Inches(
            0.55
        )

        section.left_margin = Inches(
            0.65
        )

        section.right_margin = Inches(
            0.65
        )


def apply_default_font(
    document: Document
) -> None:
    normal_style = document.styles[
        "Normal"
    ]

    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Arial"
    )


def clean_string(
    value: Any
) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).split()
    ).strip()


def clean_list(
    values: Any
) -> list[str]:
    if not isinstance(
        values,
        list
    ):
        return []

    cleaned = []
    seen = set()

    for value in values:
        clean_value = clean_string(
            value
        )

        normalized = (
            clean_value.lower()
        )

        if (
            clean_value
            and normalized not in seen
        ):
            cleaned.append(
                clean_value
            )

            seen.add(
                normalized
            )

    return cleaned


def add_header(
    document: Document,
    *,
    full_name: str,
    email: str,
    phone: str = "",
    linkedin: str = "",
    github: str = ""
) -> None:
    name_paragraph = (
        document.add_paragraph()
    )

    name_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    name_paragraph.paragraph_format.space_after = Pt(
        2
    )

    name_run = name_paragraph.add_run(
        full_name.upper()
        or "PROFESSIONAL RESUME"
    )

    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(19)
    name_run.font.color.rgb = RGBColor(
        15,
        23,
        42
    )

    contact_items = [
        clean_string(phone),
        clean_string(email),
        clean_string(linkedin),
        clean_string(github)
    ]

    contact_items = [
        item
        for item in contact_items
        if item
    ]

    if contact_items:
        contact_paragraph = (
            document.add_paragraph()
        )

        contact_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        contact_paragraph.paragraph_format.space_after = Pt(
            8
        )

        contact_run = (
            contact_paragraph.add_run(
                "  |  ".join(
                    contact_items
                )
            )
        )

        contact_run.font.name = "Arial"
        contact_run.font.size = Pt(9)
        contact_run.font.color.rgb = RGBColor(
            71,
            85,
            105
        )


def add_section_heading(
    document: Document,
    title: str
) -> None:
    paragraph = (
        document.add_paragraph()
    )

    paragraph.paragraph_format.space_before = Pt(
        7
    )

    paragraph.paragraph_format.space_after = Pt(
        4
    )

    run = paragraph.add_run(
        title.upper()
    )

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(
        15,
        118,
        110
    )

    add_bottom_border(
        paragraph
    )


def add_body_paragraph(
    document: Document,
    text: str
) -> None:
    clean_text = clean_string(
        text
    )

    if not clean_text:
        return

    paragraph = (
        document.add_paragraph()
    )

    paragraph.paragraph_format.space_after = Pt(
        3
    )

    paragraph.paragraph_format.line_spacing = 1.05

    run = paragraph.add_run(
        clean_text
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(
        30,
        41,
        59
    )


def looks_like_heading(
    value: str
) -> bool:
    clean_value = clean_string(
        value
    )

    if not clean_value:
        return False

    if len(
        clean_value.split()
    ) > 10:
        return False

    if clean_value.endswith(
        "."
    ):
        return False

    return True


def add_bullet_items(
    document: Document,
    items: list[str]
) -> None:
    safe_items = clean_list(
        items
    )

    for item in safe_items:
        if looks_like_heading(
            item
        ):
            paragraph = (
                document.add_paragraph()
            )

            paragraph.paragraph_format.space_before = Pt(
                4
            )

            paragraph.paragraph_format.space_after = Pt(
                2
            )

            run = paragraph.add_run(
                item
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(
                15,
                23,
                42
            )

            continue

        paragraph = (
            document.add_paragraph(
                style="List Bullet"
            )
        )

        paragraph.paragraph_format.left_indent = Inches(
            0.18
        )

        paragraph.paragraph_format.first_line_indent = Inches(
            -0.08
        )

        paragraph.paragraph_format.space_after = Pt(
            2
        )

        paragraph.paragraph_format.line_spacing = 1.02

        run = paragraph.add_run(
            item
        )

        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(
            30,
            41,
            59
        )


def add_skills_table(
    document: Document,
    verified_skills: list[str],
    recommended_skills: list[str]
) -> None:
    """
    Add only verified skills to the downloaded DOCX.

    Recommended skills remain visible on the website
    but are not presented as existing resume skills.
    """

    verified = clean_list(
        verified_skills
    )

    if not verified:
        return

    table = document.add_table(
        rows=0,
        cols=2
    )

    table.autofit = True

    categorized_skills = {
        "Programming": [],
        "Frontend": [],
        "Backend": [],
        "Data and AI": [],
        "Databases": [],
        "Tools and Cloud": [],
        "Other": []
    }

    programming = {
        "python",
        "java",
        "c",
        "c++",
        "c#",
        "javascript",
        "typescript",
        "sql"
    }

    frontend = {
        "html",
        "html5",
        "css",
        "css3",
        "react",
        "react.js",
        "bootstrap",
        "tailwind",
        "tailwind css"
    }

    backend = {
        "node.js",
        "nodejs",
        "express",
        "express.js",
        "fastapi",
        "flask",
        "django",
        "rest api",
        "rest apis"
    }

    data_ai = {
        "artificial intelligence",
        "machine learning",
        "deep learning",
        "pandas",
        "numpy",
        "scikit-learn",
        "sklearn",
        "natural language processing",
        "nlp",
        "power bi",
        "tableau",
        "excel",
        "data analysis",
        "data analytics",
        "data visualization",
        "eda"
    }

    databases = {
        "mongodb",
        "postgresql",
        "mysql",
        "sqlite",
        "redis",
        "sql server"
    }

    tools_cloud = {
        "git",
        "github",
        "docker",
        "aws",
        "azure",
        "gcp",
        "cloud deployment",
        "ci/cd",
        "n8n",
        "hadoop",
        "vs code",
        "postman",
        "replit"
    }

    for skill in verified:
        normalized_skill = (
            skill.lower().strip()
        )

        if normalized_skill in programming:
            category = "Programming"

        elif normalized_skill in frontend:
            category = "Frontend"

        elif normalized_skill in backend:
            category = "Backend"

        elif normalized_skill in data_ai:
            category = "Data and AI"

        elif normalized_skill in databases:
            category = "Databases"

        elif normalized_skill in tools_cloud:
            category = "Tools and Cloud"

        else:
            category = "Other"

        categorized_skills[
            category
        ].append(
            skill
        )

    for category, skills in (
        categorized_skills.items()
    ):
        if not skills:
            continue

        row = table.add_row()

        category_cell = row.cells[0]
        skills_cell = row.cells[1]

        set_cell_margins(
            category_cell
        )

        set_cell_margins(
            skills_cell
        )

        category_run = (
            category_cell
            .paragraphs[0]
            .add_run(category)
        )

        category_run.bold = True
        category_run.font.name = "Arial"
        category_run.font.size = Pt(9.5)
        category_run.font.color.rgb = RGBColor(
            15,
            118,
            110
        )

        skills_run = (
            skills_cell
            .paragraphs[0]
            .add_run(
                ", ".join(skills)
            )
        )

        skills_run.font.name = "Arial"
        skills_run.font.size = Pt(9.5)
        skills_run.font.color.rgb = RGBColor(
            30,
            41,
            59
        )


def generate_rewritten_resume_docx(
    *,
    user_name: str,
    user_email: str,
    file_name: str,
    target_career: str,
    rewritten_content: dict[str, Any],
    phone: str = "",
    linkedin: str = "",
    github: str = ""
) -> bytes:
    """
    Create an improved resume as a DOCX file and
    return the file bytes.
    """

    if not isinstance(
        rewritten_content,
        dict
    ):
        raise ValueError(
            "Rewritten resume content is invalid."
        )

    summary = clean_string(
        rewritten_content.get(
            "summary"
        )
    )

    experience = clean_list(
        rewritten_content.get(
            "experience"
        )
    )

    projects = clean_list(
        rewritten_content.get(
            "projects"
        )
    )

    verified_skills = clean_list(
        rewritten_content.get(
            "verified_skills"
        )
    )

    recommended_skills = clean_list(
        rewritten_content.get(
            "recommended_skills"
        )
    )

    document = Document()

    apply_document_margins(
        document
    )

    apply_default_font(
        document
    )

    add_header(
        document,
        full_name=clean_string(
            user_name
        ),
        email=clean_string(
            user_email
        ),
        phone=clean_string(
            phone
        ),
        linkedin=clean_string(
            linkedin
        ),
        github=clean_string(
            github
        )
    )

    if target_career:
        career_paragraph = (
            document.add_paragraph()
        )

        career_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        career_paragraph.paragraph_format.space_after = Pt(
            7
        )

        career_run = (
            career_paragraph.add_run(
                f"Target Career: {target_career}"
            )
        )

        career_run.bold = True
        career_run.font.name = "Arial"
        career_run.font.size = Pt(9.5)
        career_run.font.color.rgb = RGBColor(
            15,
            118,
            110
        )

    if summary:
        add_section_heading(
            document,
            "Professional Summary"
        )

        add_body_paragraph(
            document,
            summary
        )

    if experience:
        add_section_heading(
            document,
            "Experience"
        )

        add_bullet_items(
            document,
            experience
        )

    if projects:
        add_section_heading(
            document,
            "Projects"
        )

        add_bullet_items(
            document,
            projects
        )

    if (
        verified_skills
        or recommended_skills
    ):
        add_section_heading(
            document,
            "Technical Skills"
        )

        add_skills_table(
            document,
            verified_skills,
            recommended_skills
        )

    source_paragraph = (
        document.add_paragraph()
    )

    source_paragraph.paragraph_format.space_before = Pt(
        8
    )

    source_run = (
        source_paragraph.add_run(
            f"Generated from: {file_name}"
        )
    )

    source_run.italic = True
    source_run.font.name = "Arial"
    source_run.font.size = Pt(7.5)
    source_run.font.color.rgb = RGBColor(
        148,
        163,
        184
    )

    buffer = BytesIO()

    document.save(
        buffer
    )

    buffer.seek(0)

    return buffer.getvalue()