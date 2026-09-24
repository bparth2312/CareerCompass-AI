from pathlib import Path

import fitz
from docx import Document


def normalize_extracted_text(
    text: str
) -> str:
    """
    Clean text extracted from PDF or DOCX while
    preserving section headings and readable lines.
    """

    if not text:
        return ""

    cleaned_lines = []

    for raw_line in text.splitlines():
        line = " ".join(
            raw_line.split()
        ).strip()

        if line:
            cleaned_lines.append(line)

    return "\n".join(
        cleaned_lines
    ).strip()


def extract_text_from_pdf(
    file_path: str
) -> str:
    """
    Extract visible text and clickable hyperlinks
    from a PDF using PyMuPDF.

    Text blocks are sorted by their page position to
    improve reading order for resume layouts.
    """

    document = fitz.open(file_path)

    extracted_parts = []
    discovered_links = set()

    try:
        for page in document:
            blocks = page.get_text(
                "blocks",
                sort=True
            )

            page_lines = []

            for block in blocks:
                if len(block) < 5:
                    continue

                block_text = str(
                    block[4] or ""
                ).strip()

                if not block_text:
                    continue

                clean_block = (
                    normalize_extracted_text(
                        block_text
                    )
                )

                if clean_block:
                    page_lines.append(
                        clean_block
                    )

            if page_lines:
                extracted_parts.append(
                    "\n".join(page_lines)
                )

            for link in page.get_links():
                uri = link.get("uri")

                if not uri:
                    continue

                clean_uri = str(
                    uri
                ).strip()

                if clean_uri:
                    discovered_links.add(
                        clean_uri
                    )

    finally:
        document.close()

    if discovered_links:
        extracted_parts.append(
            "PROFILE LINKS"
        )

        extracted_parts.extend(
            sorted(discovered_links)
        )

    return "\n".join(
        extracted_parts
    ).strip()


def extract_text_from_docx(
    file_path: str
) -> str:
    """
    Extract paragraph text and hyperlink URLs from
    a DOCX resume.
    """

    document = Document(file_path)

    extracted_parts = []
    discovered_links = set()

    for paragraph in document.paragraphs:
        text = " ".join(
            paragraph.text.split()
        ).strip()

        if text:
            extracted_parts.append(text)

        relationships = (
            paragraph.part.rels
        )

        for relationship in (
            relationships.values()
        ):
            if relationship.reltype.endswith(
                "/hyperlink"
            ):
                target = str(
                    relationship.target_ref
                    or ""
                ).strip()

                if target:
                    discovered_links.add(
                        target
                    )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = " ".join(
                    cell.text.split()
                ).strip()

                if cell_text:
                    extracted_parts.append(
                        cell_text
                    )

    if discovered_links:
        extracted_parts.append(
            "PROFILE LINKS"
        )

        extracted_parts.extend(
            sorted(discovered_links)
        )

    return "\n".join(
        extracted_parts
    ).strip()


def extract_resume_text(
    file_path: str
) -> str:
    path = Path(file_path)

    if not path.exists():
        raise ValueError(
            "The resume file does not exist."
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        extracted_text = (
            extract_text_from_pdf(
                file_path
            )
        )

    elif extension == ".docx":
        extracted_text = (
            extract_text_from_docx(
                file_path
            )
        )

    else:
        raise ValueError(
            "Resume parsing is supported only "
            "for PDF and DOCX files."
        )

    if not extracted_text.strip():
        raise ValueError(
            "No readable text was found in "
            "the uploaded resume."
        )

    return extracted_text